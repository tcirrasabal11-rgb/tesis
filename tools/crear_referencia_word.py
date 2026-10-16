#!/usr/bin/env python3
"""Crea un estilo base de Word, provisional hasta recibir plantilla institucional."""
import argparse
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from tesis import output_path, ProjectError

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output',default='plantillas/referencia_base.docx')
    parser.add_argument('--overwrite',action='store_true')
    args=parser.parse_args()
    out=output_path(args.output,args.overwrite)
    doc=Document();section=doc.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Inches(1)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12)
    normal.font.color.rgb=RGBColor(0,0,0)
    normal.paragraph_format.line_spacing=1.5
    normal.paragraph_format.space_after=Pt(6)
    for name,size in [('Title',18),('Subtitle',12),('Heading 1',14),('Heading 2',12),('Heading 3',12),('Heading 4',12)]:
        style=doc.styles[name];style.font.name='Times New Roman';style.font.size=Pt(size)
        style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.space_before=Pt(12);style.paragraph_format.space_after=Pt(6)
        if name.startswith('Heading'):style.font.bold=True;style.paragraph_format.keep_with_next=True
        ppr=style.element.find(qn('w:pPr'))
        if ppr is not None:
            border=ppr.find(qn('w:pBdr'))
            if border is not None:ppr.remove(border)
    foot=section.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=foot.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
    run._r.addnext(field)
    doc.add_paragraph('Estilo base para exportación de la tesis',style='Title')
    doc.add_paragraph('Este archivo proporciona estilos y tamaño de página para la exportación con Pandoc. La plantilla institucional, cuando se reciba, determina el formato definitivo.')
    doc.add_heading('Configuración base',level=1)
    doc.add_paragraph('Página carta, márgenes de una pulgada, Times New Roman de doce puntos, interlineado de uno y medio y encabezados negros. Estos ajustes son una propuesta de trabajo que debe contrastarse con la guía de la universidad.')
    doc.add_heading('Uso',level=1)
    doc.add_paragraph('El contenido de este archivo no se incorpora al texto de la tesis. Pandoc utiliza sus estilos; los capítulos de la carpeta capitulos contienen el contenido editable.')
    doc.core_properties.title='Estilo base para exportación de la tesis'
    doc.core_properties.author=''
    doc.save(out)
    print(out)

if __name__=='__main__':
    try:main()
    except ProjectError as exc:raise SystemExit(str(exc))
